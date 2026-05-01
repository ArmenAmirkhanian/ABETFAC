import io
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import (AssessmentBatch, AssessmentDetail, Semester,
                        RubricCriterion)


SCORE_FIELDS = AssessmentDetail.SCORE_FIELDS


def build_excel_export(semester_id=None):
    """
    Build and return a BytesIO Excel workbook with all assessment data.
    Optionally filtered to a single semester.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    query = AssessmentBatch.query
    if semester_id:
        query = query.filter_by(semester_id=semester_id)
    batches = query.order_by(AssessmentBatch.semester_id, AssessmentBatch.course_id).all()

    if not batches:
        ws = wb.create_sheet('No Data')
        ws['A1'] = 'No assessment data found for the selected filter.'
        return _to_bytes(wb)

    # Summary sheet
    _build_summary_sheet(wb, batches)

    # Detail sheet
    _build_detail_sheet(wb, batches)

    # One sheet per course (if multiple batches per course, append)
    _build_course_sheets(wb, batches)

    return _to_bytes(wb)


def _hdr_style(ws, row, cols, fill_hex='1F497D'):
    fill = PatternFill('solid', fgColor=fill_hex)
    font = Font(color='FFFFFF', bold=True)
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal='center', wrap_text=True)


def _build_summary_sheet(wb, batches):
    ws = wb.create_sheet('Summary')
    headers = ['Course', 'Course Name', 'Semester', 'Faculty', 'SLO', 'Outcome',
               'Students Assessed', 'Avg Score (All Criteria)', 'Action Taken']
    for col, h in enumerate(headers, 1):
        ws.cell(1, col, h)
    _hdr_style(ws, 1, len(headers))

    for row_i, batch in enumerate(batches, 2):
        details = batch.details.all()
        if details:
            all_scores = []
            for d in details:
                for field, _ in SCORE_FIELDS:
                    s = getattr(d, field)
                    if s:
                        all_scores.append(s)
            avg = round(sum(all_scores) / len(all_scores), 2) if all_scores else ''
        else:
            avg = ''

        ws.cell(row_i, 1, batch.course_id)
        ws.cell(row_i, 2, batch.course.name)
        ws.cell(row_i, 3, batch.semester.label)
        ws.cell(row_i, 4, batch.faculty.display_name)
        ws.cell(row_i, 5, batch.slo.code)
        ws.cell(row_i, 6, batch.outcome_assessed or batch.slo.summary)
        ws.cell(row_i, 7, len(details))
        ws.cell(row_i, 8, avg)
        ws.cell(row_i, 9, batch.action_taken or '')

    _autofit(ws)


def _build_detail_sheet(wb, batches):
    ws = wb.create_sheet('All Details')
    score_labels = [label for _, label in SCORE_FIELDS]
    headers = (['Course', 'Semester', 'Faculty', 'SLO', 'Student Last', 'Student First',
                'ELEARN ID', 'Earned Grade', 'Artifact'] + score_labels)
    for col, h in enumerate(headers, 1):
        ws.cell(1, col, h)
    _hdr_style(ws, 1, len(headers))

    row_i = 2
    for batch in batches:
        for d in batch.details.order_by(AssessmentDetail.id).all():
            ws.cell(row_i, 1, batch.course_id)
            ws.cell(row_i, 2, batch.semester.label)
            ws.cell(row_i, 3, batch.faculty.display_name)
            ws.cell(row_i, 4, batch.slo.code)
            ws.cell(row_i, 5, d.student.last_name)
            ws.cell(row_i, 6, d.student.first_name)
            ws.cell(row_i, 7, d.student.elearn_id)
            ws.cell(row_i, 8, d.earned_grade or '')
            ws.cell(row_i, 9, d.artifact_filename or '')
            for col, (field, _) in enumerate(SCORE_FIELDS, 10):
                ws.cell(row_i, col, getattr(d, field) or '')
            row_i += 1

    _autofit(ws)


def _build_course_sheets(wb, batches):
    # Group by course
    by_course = {}
    for b in batches:
        by_course.setdefault(b.course_id, []).append(b)

    score_labels = [label for _, label in SCORE_FIELDS]

    for course_id, course_batches in by_course.items():
        safe_name = course_id.replace('/', '-')[:31]
        ws = wb.create_sheet(safe_name)

        ws['A1'] = f'{course_id} — {course_batches[0].course.name}'
        ws['A1'].font = Font(bold=True, size=12)
        ws.merge_cells(f'A1:{get_column_letter(9 + len(SCORE_FIELDS))}1')

        headers = (['Semester', 'SLO', 'Student Last', 'Student First',
                    'ELEARN ID', 'Grade', 'Artifact'] + score_labels + ['Avg Score'])
        for col, h in enumerate(headers, 1):
            ws.cell(2, col, h)
        _hdr_style(ws, 2, len(headers))

        row_i = 3
        for batch in course_batches:
            for d in batch.details.order_by(AssessmentDetail.id).all():
                row_scores = [getattr(d, field) for field, _ in SCORE_FIELDS]
                valid_scores = [s for s in row_scores if s]
                avg = round(sum(valid_scores) / len(valid_scores), 2) if valid_scores else ''

                ws.cell(row_i, 1, batch.semester.label)
                ws.cell(row_i, 2, batch.slo.code)
                ws.cell(row_i, 3, d.student.last_name)
                ws.cell(row_i, 4, d.student.first_name)
                ws.cell(row_i, 5, d.student.elearn_id)
                ws.cell(row_i, 6, d.earned_grade or '')
                ws.cell(row_i, 7, d.artifact_filename or '')
                for col, s in enumerate(row_scores, 8):
                    ws.cell(row_i, col, s or '')
                ws.cell(row_i, 8 + len(SCORE_FIELDS), avg)
                row_i += 1

        _autofit(ws)


def _autofit(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def _to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
